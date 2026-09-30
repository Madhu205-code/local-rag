from __future__ import annotations

import csv
import html
import io
import json
import re
from pathlib import Path
from config import WEB_EXTENSIONS

Section = tuple[str, str]

BINARY_CHECK_EXEMPT = {
    ".pdf", ".docx", ".xlsx", ".pptx", ".epub", ".html", ".htm", ".ipynb", ".png", ".jpg",
}

ENCODINGS = ("utf-8-sig", "utf-8", "utf-16", "cp1252", "latin-1")


class UnsupportedFile(Exception):
    pass


def read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ENCODINGS:
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return raw.decode("utf-8", errors="replace")


def _clean(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _by_lines(text: str, lines_per_section: int = 120) -> list[Section]:
    lines = text.split("\n")
    out: list[Section] = []
    for start in range(0, len(lines), lines_per_section):
        block = "\n".join(lines[start : start + lines_per_section]).strip()
        if block:
            out.append((f"lines {start + 1}-{min(start + lines_per_section, len(lines))}", block))
    return out


def _by_headings(text: str) -> list[Section]:
    sections: list[Section] = []
    current_label = "intro"
    buf: list[str] = []
    for line in text.split("\n"):
        if re.match(r"^#{1,6}\s+\S", line) or re.match(r"^[A-Z][\w \-/]{2,40}:$", line):
            if buf:
                sections.append((current_label, "\n".join(buf).strip()))
            current_label = line.lstrip("#").strip() or "section"
            buf = []
            continue
        buf.append(line)
    if buf:
        sections.append((current_label, "\n".join(buf).strip()))
    return [(l, b) for l, b in sections if b]


def _pdf(path: Path) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    out: list[Section] = []
    for i, page in enumerate(reader.pages):
        try:
            content = page.extract_text() or ""
        except Exception:
            content = ""
        content = _clean(content)
        if content:
            out.append((f"page {i + 1}", content))
    return out


def _docx(path: Path) -> list[Section]:
    import docx

    document = docx.Document(str(path))
    sections: list[Section] = []
    buf: list[str] = []
    label = "intro"

    def flush() -> None:
        body = "\n".join(x for x in buf if x).strip()
        if body:
            sections.append((label, body))
        buf.clear()

    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("heading") or style == "title":
            flush()
            label = text
            continue
        buf.append(text)
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [(c.text or "").strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            flush()
            label = f"table ({len(rows)} rows)"
            buf.extend(rows)
    flush()
    return sections


def _xlsx(path: Path) -> list[Section]:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out: list[Section] = []
    for sheet in workbook.worksheets:
        lines: list[str] = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = ["" if v is None else str(v).strip() for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if not cells:
                continue
            lines.append(f"r{row_index}: " + " | ".join(cells))
            if len(lines) >= 200:
                out.append((f"sheet {sheet.title} rows 1-{row_index}", "\n".join(lines)))
                lines = []
        if lines:
            out.append((f"sheet {sheet.title} ({len(lines)} rows)", "\n".join(lines)))
        sheet_label = sheet.title
    workbook.close()
    return out


def _pptx(path: Path) -> list[Section]:
    from pptx import Presentation

    deck = Presentation(str(path))
    out: list[Section] = []
    for index, slide in enumerate(deck.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if text.strip():
                    parts.append(text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [(c.text or "").strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
        body = _clean("\n".join(parts))
        if body:
            out.append((f"slide {index}", body))
    return out


def _epub(path: Path) -> list[Section]:
    import ebooklib
    from ebooklib import epub

    book = epub.read_epub(str(path), options={"ignore_ncx": True})
    out: list[Section] = []
    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
        raw = item.get_content()
        try:
            text = raw.decode("utf-8", errors="replace")
        except AttributeError:
            text = str(raw)
        text = re.sub(r"<[^>]+>", " ", text)
        text = _clean(html.unescape(text))
        if len(text) > 30:
            out.append((item.get_name() or "chapter", text))
    return out


def _html(path: Path) -> list[Section]:
    from bs4 import BeautifulSoup

    raw = read_text_file(path)
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "head"]):
        tag.decompose()
    title = soup.title.string.strip() if soup.title and soup.title.string else path.name
    body = _clean(soup.get_text("\n"))
    if not body:
        return []
    body = f"Page title: {title}\n{body}"
    return _by_lines(body, 150)


def _notebook(path: Path) -> list[Section]:
    payload = json.loads(read_text_file(path))
    cells = payload.get("cells", []) if isinstance(payload, dict) else []
    out: list[Section] = []
    for index, cell in enumerate(cells):
        source = cell.get("source", [])
        text = "".join(source) if isinstance(source, list) else str(source)
        text = text.strip()
        if not text:
            continue
        kind = cell.get("cell_type", "code")
        out.append((f"cell {index + 1} [{kind}]", text))
    return out


def _jsonish(path: Path) -> list[Section]:
    raw = read_text_file(path)
    try:
        payload = json.loads(raw)
    except json.JSONDecodeError:
        return _by_lines(raw)

    lines: list[str] = []

    def walk(node, trail: str = "", depth: int = 0) -> None:
        if len(lines) >= 4000 or depth > 8:
            return
        if isinstance(node, dict):
            for key, value in node.items():
                walk(value, f"{trail}.{key}" if trail else str(key), depth + 1)
        elif isinstance(node, list):
            for idx, value in enumerate(node[:400]):
                walk(value, f"{trail}[{idx}]", depth + 1)
        elif node is None:
            return
        else:
            lines.append(f"{trail}: {node}")

    walk(payload)
    body = _clean("\n".join(lines))
    return [("json values", body)] if body else []


def _delimited(path: Path) -> list[Section]:
    raw = read_text_file(path)
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        dialect = csv.Sniffer().sniff(raw[:8000], delimiters=",\t;|")
        delimiter = dialect.delimiter
    except csv.Error:
        pass
    rows: list[str] = []
    out: list[Section] = []
    reader = csv.reader(io.StringIO(raw), delimiter=delimiter)
    header = next(reader, None)
    if header:
        rows.append("columns: " + " | ".join(str(h) for h in header))
    for index, row in enumerate(reader, start=1):
        cells = [str(c).strip() for c in row]
        if any(cells):
            rows.append(f"row {index}: " + " | ".join(cells))
        if len(rows) >= 200:
            out.append((f"rows 1-{index}", "\n".join(rows)))
            rows = []
    if rows:
        out.append((f"rows ({len(rows)})", "\n".join(rows)))
    return out


def _rtf(path: Path) -> list[Section]:
    raw = read_text_file(path)
    text = re.sub(r"\\par[d]?\b", "\n", raw)
    text = re.sub(r"\{\\\*[^{}]*\}", "", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d*\s?", "", text)
    text = text.replace("{", "").replace("}", "")
    text = _clean(html.unescape(text))
    return _by_lines(text, 150) if text else []


def load_sections(path: Path) -> tuple[str, list[Section]]:
    suffix = path.suffix.lower()
    title = path.stem.replace("_", " ").replace("-", " ").strip() or path.name

    if suffix == ".pdf":
        return title, _pdf(path)
    if suffix == ".docx":
        return title, _docx(path)
    if suffix == ".xlsx":
        return title, _xlsx(path)
    if suffix == ".pptx":
        return title, _pptx(path)
    if suffix == ".epub":
        return title, _epub(path)
    if suffix == ".rtf":
        return title, _rtf(path)
    if suffix == ".ipynb":
        return title, _notebook(path)
    if suffix in WEB_EXTENSIONS:
        sections = _html(path)
        if sections:
            first = sections[0][1].split("\n", 1)[0]
            if first.lower().startswith("page title:"):
                title = first.split(":", 1)[1].strip() or title
        return title, sections
    if suffix == ".json" or suffix.endswith("jsonl") or suffix.endswith("ndjson"):
        return title, _jsonish(path)
    if suffix in {".csv", ".tsv"}:
        return title, _delimited(path)
    if suffix == ".bat" or suffix == ".cmd" or suffix.endswith(".ps1"):
        return title, _by_lines(_clean(read_text_file(path)), 150)

    text = _clean(read_text_file(path))
    if suffix in {".md", ".markdown", ".mdx", ".rst", ".adoc"}:
        sections = _by_headings(text)
        return title, sections or _by_lines(text)
    return title, _by_lines(text)


def is_probably_binary(path: Path) -> bool:
    if path.suffix.lower() in BINARY_CHECK_EXEMPT:
        return False
    try:
        with open(path, "rb") as handle:
            sample = handle.read(4096)
    except OSError:
        return True
    if not sample:
        return True
    if b"\x00" in sample:
        return True
    printable = sum(1 for byte in sample if 32 <= byte <= 126 or byte in (9, 10, 13))
    return printable / max(len(sample), 1) < 0.85
